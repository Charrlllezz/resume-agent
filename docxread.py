#!/usr/bin/env python3
"""Read a DOCX's typography, colour and layout. Same shape as pdfread.read().

A DOCX is the easier document to read, not the harder one. It stores structure
rather than glyphs at coordinates, so its text comes out with no ligatures to
undo, no columns to un-interleave, and dates still attached to the role they
belong to. Everything a PDF made us infer is stated here: the run's font, its
size, its colour, the paragraph's alignment, whether the section has two
columns, whether the layout is a table with a shaded cell down one side.

It only looked worse because none of it was being read.
"""

import collections
import re

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def clean_font(name: str) -> str:
    name = re.split(r"[-,_]", (name or "").strip())[0]
    return re.sub(r"[^a-z0-9]", "", name.lower())


def _run_font(run, para, default_name, default_size):
    name = run.font.name
    size = run.font.size.pt if run.font.size else None
    style = getattr(para, "style", None)
    if not name and style is not None and style.font.name:
        name = style.font.name
    if size is None and style is not None and style.font.size:
        size = style.font.size.pt
    return clean_font(name or default_name), round(float(size or default_size), 1)


def _run_colour(run, para) -> str:
    colour = run.font.color
    if colour is not None and colour.rgb is not None:
        return "#" + str(colour.rgb).upper()
    style = getattr(para, "style", None)
    if style is not None and style.font.color is not None and style.font.color.rgb is not None:
        return "#" + str(style.font.color.rgb).upper()
    # Unset means the document default, which is black unless a theme says
    # otherwise. Counting it as "unknown" threw away most of the page.
    return "#000000"


def _shaded_column(document):
    """A table cell with a fill down one side: Word's usual sidebar.

    Word has real two-column sections, but resume templates overwhelmingly use
    a two-cell table instead, because it lets the columns run to different
    lengths.
    """
    for table in document.tables:
        if not table.rows or len(table.columns) != 2:
            continue
        for index, cell in enumerate(table.rows[0].cells[:2]):
            shd = cell._tc.xpath(".//w:shd")
            fill = shd[0].get(f"{W}fill") if shd else None
            if not fill or fill.lower() in ("auto", "ffffff", "none"):
                continue
            widths = []
            for c in table.rows[0].cells[:2]:
                tc_w = c._tc.xpath("./w:tcPr/w:tcW")
                widths.append(float(tc_w[0].get(f"{W}w") or 0) if tc_w else 0.0)
            total = sum(widths) or 1.0
            return {"side": "left" if index == 0 else "right",
                    "width_pct": max(22, min(45, round(100 * widths[index] / total))),
                    "colour": "#" + fill.upper(), "table": table, "index": index}
    return None


KNOWN_SECTIONS = {
    "contact": "contact", "details": "contact", "info": "contact",
    "skills": "skills", "expertise": "skills", "competencies": "skills",
    "education": "education", "academic": "education",
    "certifications": "certifications", "certificates": "certifications",
    "licenses": "certifications",
}


def read(data: bytes) -> dict:
    """{fonts, colours, panel, sidebar_ink, title_align, sections} or {}."""
    import io

    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        return {}

    normal = document.styles["Normal"].font
    default_name = normal.name or "Calibri"
    default_size = normal.size.pt if normal.size else 11.0

    fonts = collections.Counter()
    colours = collections.Counter()
    biggest = (0.0, None)          # (size, paragraph)

    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    for para in paragraphs:
        for run in para.runs:
            n = len(run.text.strip())
            if not n:
                continue
            family, size = _run_font(run, para, default_name, default_size)
            if family:
                fonts[(family, size)] += n
            colours[_run_colour(run, para)] += n
            if size > biggest[0]:
                biggest = (size, para)

    if not fonts:
        return {}

    panel = _shaded_column(document)
    sections = {}
    sidebar_ink = ""
    if panel:
        # The text colour inside the panel, measured. For a PDF this needs an
        # x-band; here the cell is the band.
        inside = collections.Counter()
        for para in panel["table"].rows[0].cells[panel["index"]].paragraphs:
            for run in para.runs:
                n = len(run.text.strip())
                if n:
                    inside[_run_colour(run, para)] += n
        if inside:
            sidebar_ink = inside.most_common(1)[0][0]
    if panel:
        for index, cell in enumerate(panel["table"].rows[0].cells[:2]):
            column = "side" if index == panel["index"] else "main"
            for para in cell.paragraphs:
                text = para.text.strip(" :")
                if len(text) < 4 or not text.isupper():
                    continue
                name = KNOWN_SECTIONS.get(re.sub(r"[^a-z]", "", text.lower()))
                if name:
                    sections.setdefault(name, column)

    align = biggest[1].alignment if biggest[1] is not None else None
    return {
        "page": (612.0, 792.0),
        "fonts": [(f, s, n) for (f, s), n in fonts.most_common()],
        "colours": colours.most_common(),
        "panel": ({"x0": 0.0, "width": panel["width_pct"] / 100 * 612.0,
                   "colour": panel["colour"], "side": panel["side"]} if panel else None),
        "sidebar_ink": sidebar_ink,
        "title_align": "center" if align is not None and "CENTER" in str(align) else "left",
        "sections": sections,
    }
